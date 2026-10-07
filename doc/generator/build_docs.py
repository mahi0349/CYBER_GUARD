"""
Complete 36-Section Documentation Builder for QuantumVault
Generates both .docx and .md versions in doc/ directory.
"""
import os
import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOC_DIR = Path("doc")
DOC_DIR.mkdir(exist_ok=True)
DOCX_PATH = DOC_DIR / "Quantum_Vault_Complete_Documentation.docx"
MD_PATH = DOC_DIR / "Quantum_Vault_Complete_Documentation.md"

# Style Constants
COLOR_NAVY = RGBColor(15, 23, 42)      # #0F172A
COLOR_BLUE = RGBColor(2, 132, 199)     # #0284C7
COLOR_SKY = RGBColor(14, 165, 233)     # #0EA5E9
COLOR_GRAY = RGBColor(100, 116, 139)   # #64748B
COLOR_BODY = RGBColor(51, 65, 85)      # #334155
COLOR_WHITE = RGBColor(255, 255, 255)
HEX_NAVY = "0F172A"
HEX_BLUE = "0284C7"
HEX_LIGHT_ROW = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BG = "F0F9FF"
HEX_CODE_BG = "F1F5F9"

class DocBuilder:
    def __init__(self):
        self.doc = docx.Document()
        self.md_lines = []
        self._setup_page()

    def _setup_page(self):
        for s in self.doc.sections:
            s.top_margin = Inches(1.0)
            s.bottom_margin = Inches(1.0)
            s.left_margin = Inches(1.0)
            s.right_margin = Inches(1.0)
            
            # Header
            header = s.header
            hp = header.paragraphs[0]
            hp.text = "Quantum Vault — Complete Software Documentation"
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if hp.runs:
                hp.runs[0].font.name = "Segoe UI"
                hp.runs[0].font.size = Pt(8.5)
                hp.runs[0].font.color.rgb = COLOR_GRAY

            # Footer
            footer = s.footer
            fp = footer.paragraphs[0]
            fp.text = "Enterprise AI Cyber Threat Detection & Endpoint Defense  •  v1.0.0"
            fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if fp.runs:
                fp.runs[0].font.name = "Segoe UI"
                fp.runs[0].font.size = Pt(8.5)
                fp.runs[0].font.color.rgb = COLOR_GRAY

    def add_title(self, text, subtitle=None):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(36)
        p.paragraph_format.space_after = Pt(12)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(28)
        r.font.bold = True
        r.font.color.rgb = COLOR_NAVY

        if subtitle:
            ps = self.doc.add_paragraph()
            ps.paragraph_format.space_after = Pt(24)
            ps.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rs = ps.add_run(subtitle)
            rs.font.name = "Segoe UI"
            rs.font.size = Pt(13)
            rs.font.color.rgb = COLOR_BLUE

        self.md_lines.append(f"# {text}\n")
        if subtitle:
            self.md_lines.append(f"### {subtitle}\n")

    def add_h1(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = COLOR_NAVY

        self.md_lines.append(f"\n# {text}\n")

    def add_h2(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = COLOR_BLUE

        self.md_lines.append(f"\n## {text}\n")

    def add_h3(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = COLOR_NAVY

        self.md_lines.append(f"\n### {text}\n")

    def add_p(self, text, bold_prefix=None, italic=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rp = p.add_run(bold_prefix)
            rp.font.name = "Segoe UI"
            rp.font.size = Pt(10)
            rp.font.bold = True
            rp.font.color.rgb = COLOR_NAVY
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(10)
        r.font.italic = italic
        r.font.color.rgb = COLOR_BODY

        full_md = (f"**{bold_prefix}** " if bold_prefix else "") + (f"*{text}*" if italic else text)
        self.md_lines.append(f"{full_md}\n")

    def add_bullet(self, bold_prefix, text):
        p = self.doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rp = p.add_run(bold_prefix)
            rp.font.name = "Segoe UI"
            rp.font.size = Pt(10)
            rp.font.bold = True
            rp.font.color.rgb = COLOR_NAVY
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_BODY

        self.md_lines.append(f"- **{bold_prefix}** {text}")

    def add_callout(self, text, title="NOTE"):
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        
        # XML styling
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_CALLOUT_BG}"/>')
        cell._tc.get_or_add_tcPr().append(shd)
        
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_BLUE}"/>'
            f'  <w:top w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        cell._tc.get_or_add_tcPr().append(borders)
        
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'  <w:top w:w="140" w:type="dxa"/><w:bottom w:w="140" w:type="dxa"/>'
            f'  <w:left w:w="200" w:type="dxa"/><w:right w:w="200" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        cell._tc.get_or_add_tcPr().append(tcMar)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        rt = p.add_run(f"[{title}] ")
        rt.font.name = "Segoe UI"
        rt.font.size = Pt(9.5)
        rt.font.bold = True
        rt.font.color.rgb = COLOR_BLUE
        
        rx = p.add_run(text)
        rx.font.name = "Segoe UI"
        rx.font.size = Pt(9.5)
        rx.font.color.rgb = COLOR_BODY
        
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        self.md_lines.append(f"\n> **[{title}]**: {text}\n")

    def add_code(self, code_text):
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_CODE_BG}"/>')
        cell._tc.get_or_add_tcPr().append(shd)
        
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'</w:tcBorders>'
        )
        cell._tc.get_or_add_tcPr().append(borders)
        
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'  <w:top w:w="120" w:type="dxa"/><w:bottom w:w="120" w:type="dxa"/>'
            f'  <w:left w:w="160" w:type="dxa"/><w:right w:w="160" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        cell._tc.get_or_add_tcPr().append(tcMar)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(code_text)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(30, 41, 59)
        
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        self.md_lines.append(f"\n```\n{code_text}\n```\n")

    def add_table(self, headers, data, col_widths=None):
        tbl = self.doc.add_table(rows=len(data) + 1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        tblPr = tbl._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
            f'  <w:insideV w:val="none"/><w:left w:val="none"/><w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)
        
        hdr_row = tbl.rows[0]
        hdr_row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        
        for idx, title in enumerate(headers):
            cell = hdr_row.cells[idx]
            cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_NAVY}"/>'))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(title)
            r.font.name = "Segoe UI"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = COLOR_WHITE
            
        for r_idx, row_data in enumerate(data):
            row = tbl.rows[r_idx + 1]
            bg = HEX_LIGHT_ROW if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row_data):
                cell = row.cells[c_idx]
                if bg != "FFFFFF":
                    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>'))
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                r = p.add_run(str(val))
                r.font.name = "Segoe UI"
                r.font.size = Pt(9)
                r.font.color.rgb = COLOR_BODY
                
        if col_widths:
            for row in tbl.rows:
                for idx, w in enumerate(col_widths):
                    if idx < len(row.cells):
                        row.cells[idx].width = Inches(w)
                        
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # Markdown Table
        md_hdr = "| " + " | ".join(headers) + " |"
        md_sep = "| " + " | ".join(["---"] * len(headers)) + " |"
        md_rows = ["| " + " | ".join(str(c).replace("\n", " ") for c in row) + " |" for row in data]
        self.md_lines.append("\n" + "\n".join([md_hdr, md_sep] + md_rows) + "\n")

    def page_break(self):
        self.doc.add_page_break()
        self.md_lines.append("\n---\n")

    def save(self):
        self.doc.save(str(DOCX_PATH))
        with open(MD_PATH, "w", encoding="utf-8") as f:
            f.write("\n".join(self.md_lines))
        print(f"[+] Saved Word documentation: {DOCX_PATH} ({DOCX_PATH.stat().st_size} bytes)")
        print(f"[+] Saved Markdown documentation: {MD_PATH} ({MD_PATH.stat().st_size} bytes)")

print("DocBuilder framework initialized.")
